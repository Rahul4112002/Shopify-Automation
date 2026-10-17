from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from catalog_rules import (
    clean_ai_name,
    compose_title,
    derive_style_name,
    normalize_color,
    normalize_product_type,
    normalize_size,
)
from processors import (
    fill_product_data,
    load_gs1,
    map_product_images,
    normalize_dropbox_url,
    run_pipeline,
    transform_content_master,
)


def test_content_transform_inserts_generated_columns():
    df = pd.DataFrame(
        {
            "BZ CODE": ["SKU1"],
            "Product Title": ["Lavie Women's Sling - Black"],
            "Colour": ["p blue"],
            "Bullet Point 1": ["Roomy"],
            "Bullet Point 2": ["Light"],
            "Special Feature 1 ": ["Zip closure."],
            "Brand Name": ["lavie"],
        }
    )

    result, info = transform_content_master(df)

    assert result.loc[0, "Title"] == "Lavie Sling Bag Black"
    assert result.loc[0, "Colour"] == "Pale Blue"
    assert result.loc[0, "Bullet Points HTML"] == "<li>Roomy</li><li>Light</li>"
    assert result.loc[0, "HTML content"] == (
        "<ul><li>Roomy</li><li>Light</li>"
        "<li><b>Special Features : </b>Zip closure</li></ul>"
    )
    assert info["product_title_col"] == "Product Title"


def test_load_gs1_detects_shifted_header(tmp_path: Path):
    path = tmp_path / "gs1.xlsx"
    raw = pd.DataFrame(
        [
            ["ignore", "ignore", "ignore"],
            ["Article", "EAN/UPC", "NEW MRP"],
            ["SKU1", "123456", 999],
        ]
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        raw.to_excel(writer, index=False, header=False, sheet_name="Master")

    df, info = load_gs1(path)

    assert info.header_row == 1
    assert info.sheet_name == "Master"
    assert list(df.columns) == ["Article", "EAN/UPC", "NEW MRP"]
    assert df.loc[0, "Article"] == "SKU1"


def test_load_gs1_prefers_master_over_discontinued_sheet(tmp_path: Path):
    path = tmp_path / "gs1.xlsx"
    # The real GS1 master has a blank first column; the discontinued sheet does not.
    master = pd.DataFrame(
        [
            [None, None, None, None, None, None, None],
            [None, "Article", "EAN/UPC", "Colour", "Size", "NEW MRP", "Country "],
            ["x", "NEW1", 8903606612252, "RED", "SMALL", 6499, "CHINA"],
        ]
    )
    discontinued = pd.DataFrame(
        [
            [None, None, None, None, None, None],
            ["Article", "EAN/UPC", "Color", "Size ", "NEW MRP", "Country "],
            ["OLD1", 8903606400000, "BLACK", "LARGE", 999, "INDIA"],
        ]
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        master.to_excel(writer, index=False, header=False, sheet_name="Master")
        discontinued.to_excel(writer, index=False, header=False, sheet_name="Discontinued & Blocked")

    df, info = load_gs1(path)

    assert info.sheet_name == "Master"
    assert info.header_row == 1
    assert df.loc[0, "Article"] == "NEW1"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Satchel", "Satchel Bag"),
        ("Satchel ", "Satchel Bag"),
        ("SATCHEL", "Satchel Bag"),
        ("Sathcel Bag", "Satchel Bag"),
        ("satchel bag ", "Satchel Bag"),
        ("Tote", "Tote Bag"),
        ("tote", "Tote Bag"),
        ("SLING BAG", "Sling Bag"),
        ("Slingbag", "Sling Bag"),
        ("Sling Bags", "Sling Bag"),
        ("Party Slings", "Sling Bag"),
        ("Laptop Hangbag", "Laptop Handbag"),
        ("Laptop handbag ", "Laptop Handbag"),
        ("WALLET", "Wallet"),
        ("Framed Clutch", "Clutch"),
        ("Frame bag", "Frame Bag"),
        ("Hobo", "Hobo Bag"),
        ("Laptop Backpack", "Laptop Backpack"),
        ("School Backpack", "Backpack"),
        ("Travel Wheel Duffle", "Wheel Duffle"),
        ("women Giftset", "Gift Set"),
        ("Perfume EDP", "Perfume"),
        ("Trunk Bag", "Trunk Bag"),
        (" Backpack", "Backpack"),
    ],
)
def test_subcategory_type_normalisation(raw, expected):
    product_type, _method = normalize_product_type(raw)
    assert product_type == expected


def test_type_uses_product_category_when_subcategory_is_a_sentence():
    product_type, method = normalize_product_type(
        "A sleek essential for those who like to travel light.", "wallet"
    )
    assert (product_type, method) == ("Wallet", "category")


def test_type_keeps_unknown_subcategory_but_flags_it():
    product_type, method = normalize_product_type("Lunch Kit")
    assert (product_type, method) == ("Lunch Kit", "unmapped")


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("BLACK", "Black"),
        ("OFF WHITE/TAN", "Off White & Tan"),
        ("OFFWHITE", "Off White"),
        ("BLACK/BLACK", "Black"),
        ("LT PINK", "Light Pink"),
        ("DK.GREY", "Dark Grey"),
        ("P BLUE", "Pale Blue"),
        ("MT.GOLD", "Metallic Gold"),
        ("BK+NY+TL", "Black & Navy & Teal"),
        ("OCHER", "Ochre"),
    ],
)
def test_gs_colour_normalisation(raw, expected):
    assert normalize_color(raw) == expected


def test_size_normalisation():
    assert normalize_size("X-LARGE") == "Extra Large"
    assert normalize_size("MEDIUM") == "Medium"
    assert normalize_size("25ML") == "25 ml"


def test_lavie_rule_based_name_and_title():
    name = derive_style_name(
        product_title="Lavie Signature Whitney 2c Women's Satchel Bag (Offwhite)",
        bz_description="LS WHITNEY 2C SAT - LADIES HANDBAG",
        brand="LAVIE SIGNATURE",
        color="Off White",
        size="Medium",
        product_type="Satchel Bag",
    )
    assert name == "Whitney 2 Compartment"
    assert compose_title(name, "Medium", "Off White", "Satchel Bag") == (
        "Whitney 2 Compartment Medium Off White Satchel Bag"
    )


def test_rule_name_uses_sap_text_when_title_is_generic():
    name = derive_style_name(
        product_title="Lavie Paris Women's Soft lg Tote  (Tan)",
        bz_description="SOFT STEPH VT TOTE-LADIES HANDBAG",
        color="Tan",
        size="Medium",
        product_type="Satchel Bag",
    )
    assert name == "Soft Steph Vertical"


def test_compose_title_has_no_redundant_words():
    assert compose_title("Glare Small", "Small", "Wine", "Frame Bag") == "Glare Small Wine Frame Bag"
    assert compose_title("Canva Combo Set Tote Bag", "Large", "Wine", "Tote Bag") == (
        "Canva Combo Set Large Wine Tote Bag"
    )
    assert compose_title("Duo Quad Tan", "Large", "Tan", "Handbag") == "Duo Quad Large Tan Handbag"


def test_clean_ai_name_removes_brand_and_size():
    assert clean_ai_name('"Lavie Luxe Glare Small"') == "Glare"
    assert clean_ai_name("") == ""
    assert clean_ai_name("x " * 20) == ""


def _content_and_gs1():
    content = pd.DataFrame(
        {
            "BZ CODE": ["SKU1", "SKU2"],
            "Product Title": [
                "Lavie Sport Clean Women's Backpack (Black)",
                "Lavie Sport Clean Women's Backpack (Black)",
            ],
            "HTML content": ["<ul><li>One</li></ul>", "<ul><li>Two</li></ul>"],
            "Brand Name": ["lavie sport", "lavie sport"],
            "Colour": ["Grey", "Grey"],
            "Product Category": ["backpack", "backpack"],
            "Subcategory": ["School backpack ", "School backpack "],
            "Product Description": ["Good bag", "Good bag 2"],
            "Target Gender": ["Female", "Female"],
            "Closure type": ["Zip", "Zip"],
        }
    )
    gs1 = pd.DataFrame(
        {
            "Article": ["SKU1", "SKU2"],
            "Colour": ["BLACK/BLACK", "BLACK"],
            "Size": ["x-large", "x-large"],
            "Country": ["India", "India"],
            "NEW MRP": [1999.0, 1999.0],
            "EAN/UPC": ["8.900000000000E+12", "8900000000001.0"],
            "Family": ["CLEAN", "CLEAN"],
            "System Dimension": ["LARGE;(40L X 9W X 26H)cm", "LARGE;(40L X 9W X 26H)cm"],
        }
    )
    return content, gs1


def _product_template(skus):
    columns = [
        "SKU", "Handle", "Title", "Body (HTML)", "Vendor", "Type", "Tags", "Option1 Value",
        "Option2 Value", "Variant SKU", "Variant Price", "Variant Compare At Price",
        "Variant Barcodes", "Google Shopping / MPN", "SEO Title", "Image Alt Text", "Status",
        "Badges (product.metafields.custom.badges)",
        "product category (product.metafields.custom.product_category)",
        "Color (product.metafields.shopify.color-pattern)",
        "Size (product.metafields.shopify.size)",
        "Target gender (product.metafields.shopify.target-gender)",
        "Dimensions (product.metafields.my_fields.specifications)",
        "Manufacturer Details (product.metafields.my_fields.manufacturer_details)",
    ]
    data = {col: [""] * len(skus) for col in columns}
    data["SKU"] = list(skus)
    data["Color (product.metafields.shopify.color-pattern)"] = ["Should Clear"] * len(skus)
    return pd.DataFrame(data)


def test_product_fill_maps_content_and_gs1_fields_for_lavie():
    content, gs1 = _content_and_gs1()
    result, stats = fill_product_data(content, gs1, _product_template(["SKU1", "SKU2"]), store="lavie")

    assert stats["cm_matches"] == 2
    assert stats["gs1_matches"] == 2
    assert result.loc[0, "Title"] == "Clean Extra Large Black Backpack"
    assert result.loc[0, "Handle"] == "lavie-sport-clean-extra-large-black-backpack"
    assert result.loc[1, "Handle"] == "lavie-sport-clean-extra-large-black-backpack-1"
    assert result.loc[0, "Body (HTML)"] == "<ul><li>One</li></ul>"
    assert result.loc[0, "Vendor"] == "Lavie Sport"
    # Type comes from the Subcategory, cleaned to the store type list.
    assert result.loc[0, "Type"] == "Backpack"
    assert result.loc[0, "product category (product.metafields.custom.product_category)"] == "Backpack"
    # Option1 Value comes from the GS colour, not the Content Master colour.
    assert result.loc[0, "Option1 Value"] == "Black"
    assert result.loc[0, "Option2 Value"] == "Extra Large"
    assert result.loc[0, "Tags"] == "black, nano, new-launch, extra large"
    assert result.loc[0, "Variant SKU"] == "SKU1"
    assert result.loc[0, "Variant Price"] == 1999
    assert result.loc[0, "Variant Barcodes"] == "8900000000000"
    assert result.loc[1, "Variant Barcodes"] == "8900000000001"
    assert result.loc[0, "Google Shopping / MPN"] == "8900000000000"
    assert result.loc[0, "SEO Title"] == "Buy Clean Extra Large Black Backpack Online - Lavie World"
    assert result.loc[0, "Image Alt Text"] == result.loc[0, "Title"]
    assert result.loc[0, "Badges (product.metafields.custom.badges)"] == "New"
    assert result.loc[0, "Color (product.metafields.shopify.color-pattern)"] == ""
    assert result.loc[0, "Status"] == "draft"
    assert result.loc[0, "Dimensions (product.metafields.my_fields.specifications)"] == "40L X 9W X 26H cm"
    assert "Manufactured by" in result.loc[0, "Manufacturer Details (product.metafields.my_fields.manufacturer_details)"]
    assert stats["title_info"]["rule_titles"] == 2


def test_akiki_titles_use_gs_family():
    content, gs1 = _content_and_gs1()
    content["Brand Name"] = "Akiki London"
    content["Subcategory"] = "Satchel"
    gs1["Family"] = ["MALDON", "MALDON"]
    gs1["Size"] = ["SMALL", "SMALL"]
    gs1["Colour"] = ["RED", "MT. PINK"]

    result, stats = fill_product_data(content, gs1, _product_template(["SKU1", "SKU2"]), store="akiki")

    assert result.loc[0, "Title"] == "Maldon Small Red Satchel Bag"
    assert result.loc[1, "Title"] == "Maldon Small Metallic Pink Satchel Bag"
    assert result.loc[0, "Vendor"] == "Akiki"
    assert result.loc[0, "Handle"] == "akiki-london-maldon-small-red-satchel-bag"
    assert result.loc[0, "SEO Title"] == "Buy Maldon Small Red Satchel Bag Online - Akiki London"
    assert result.loc[0, "Image Alt Text"] == "Maldon Small Red Satchel Bag"
    assert stats["title_info"]["family_titles"] == 2


def test_lavie_ai_titles_are_shared_per_style_and_fall_back_on_bad_output():
    content, gs1 = _content_and_gs1()
    content.loc[1, "Product Title"] = "Lavie Sport Other Women's Backpack (Black)"
    calls = []

    def fake_ai(items):
        calls.append(items)
        names = {}
        for item in items:
            if "Clean" in item["product_title"]:
                names[item["id"]] = "Clean Pro Max"
            else:
                names[item["id"]] = "(" * 3  # unusable answer -> rule-based fallback
        return names

    result, stats = fill_product_data(
        content, gs1, _product_template(["SKU1", "SKU2"]), store="lavie", title_generator=fake_ai
    )

    assert len(calls) == 1
    assert len(calls[0]) == 2  # one item per style
    assert result.loc[0, "Title"] == "Clean Pro Max Extra Large Black Backpack"
    assert result.loc[1, "Title"] == "Other Extra Large Black Backpack"
    assert stats["title_info"]["ai_titles"] == 1
    assert stats["title_info"]["rule_titles"] == 1


def test_dropbox_mapping_is_headerless_and_keeps_more_than_ten_images(tmp_path: Path):
    dropbox_path = tmp_path / "dropbox.xlsx"
    many_images = [f"https://www.dropbox.com/scl/fi/file/SKU1_{idx}.jpg?dl=0" for idx in range(1, 12)]
    sport = pd.DataFrame([["SKU1", *many_images]])
    other = pd.DataFrame([["OTHER", "https://example.com/other.jpg"]])
    with pd.ExcelWriter(dropbox_path, engine="openpyxl") as writer:
        other.to_excel(writer, index=False, header=False, sheet_name="Other")
        sport.to_excel(writer, index=False, header=False, sheet_name="Sport")

    product = pd.DataFrame(
        {
            "SKU": ["SKU1", "", "SKU2"],
            "Handle": ["h1", "", "h2"],
            "Image Src": ["", "", ""],
            "Image Position": ["", "", ""],
            "Variant Image": ["", "", ""],
            "Image Alt Text": ["Alt 1", "", "Alt 2"],
        }
    )

    result, stats = map_product_images(product, dropbox_path)

    assert stats["selected_sheet"] == "Sport"
    assert stats["skus_with_images"] == 1
    assert stats["blank_sku_rows_preserved"] == 1
    assert stats["missing_image_skus"] == ["SKU2"]
    assert result.iloc[0]["Image Position"] == "1"
    assert "dl.dropboxusercontent.com" in result.iloc[0]["Image Src"]
    assert result.iloc[0]["Variant Image"] == result.iloc[0]["Image Src"]
    assert result.iloc[1]["Variant Image"] == ""
    assert result[result["SKU"] == "SKU1"].shape[0] == 11
    assert result[result["SKU"] == "SKU2"].shape[0] == 1
    assert result[result["SKU"] == ""].shape[0] == 1
    image_rows = result[(result["SKU"] == "SKU1") & (result["Image Src"] != "")]
    assert image_rows["Handle"].eq("h1").all()
    assert image_rows["Image Alt Text"].eq("Alt 1").all()


def test_dropbox_is_optional():
    product = pd.DataFrame({"SKU": ["SKU1", ""], "Image Src": ["", ""], "Handle": ["h1", ""]})

    result, stats = map_product_images(product, None)

    assert stats["skipped"] is True
    assert stats["missing_image_skus"] == []
    assert len(result) == 2
    assert result.iloc[0]["Image Src"] == ""


def test_direct_dropbox_urls_are_preserved():
    url = "https://dl.dropboxusercontent.com/scl/fi/file/image.jpg?dl=0"
    assert normalize_dropbox_url(url) == url


def test_run_pipeline_writes_numeric_barcodes(tmp_path: Path):
    content, gs1 = _content_and_gs1()
    content_path = tmp_path / "content.xlsx"
    gs1_path = tmp_path / "gs1.xlsx"
    product_path = tmp_path / "products.csv"
    content.to_excel(content_path, index=False, sheet_name="New Format")
    with pd.ExcelWriter(gs1_path, engine="openpyxl") as writer:
        gs1.to_excel(writer, index=False, sheet_name="Master")
    _product_template(["SKU1", "SKU2"]).to_csv(product_path, index=False)

    result = run_pipeline(content_path, gs1_path, None, product_path, tmp_path / "out", store="lavie")

    csv_text = (tmp_path / "out" / result["downloads"]["csv"]).read_text(encoding="utf-8-sig")
    assert "8900000000000" in csv_text
    assert "E+12" not in csv_text
    assert "8900000000000.0" not in csv_text

    workbook = openpyxl.load_workbook(tmp_path / "out" / result["downloads"]["excel"])
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    barcode_cell = sheet.cell(row=2, column=headers.index("Variant Barcodes") + 1)
    mpn_cell = sheet.cell(row=2, column=headers.index("Google Shopping / MPN") + 1)
    assert barcode_cell.value == 8900000000000
    assert barcode_cell.number_format == "0"
    assert mpn_cell.value == 8900000000000
    assert result["image_stats"]["skipped"] is True
    assert (tmp_path / "out" / result["downloads"]["report"]).exists()
